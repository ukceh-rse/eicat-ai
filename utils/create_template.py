from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_LOOKUPS_SHEET = "Lookups"


def _add_validation(ws, wb: Workbook, col_letter: str, col_def: dict) -> None:
    values = col_def["values"]
    name = col_def["name"]

    if any("," in str(v) for v in values):
        # Values contain commas — store them in a hidden sheet and reference by range,
        # because Excel's comma-delimited inline formula can't handle commas in values.
        if _LOOKUPS_SHEET not in wb.sheetnames:
            lws = wb.create_sheet(_LOOKUPS_SHEET)
            lws.sheet_state = "hidden"
            next_col = 1
        else:
            lws = wb[_LOOKUPS_SHEET]
            next_col = lws.max_column + 1

        for row_idx, v in enumerate(values, start=1):
            lws.cell(row=row_idx, column=next_col, value=v)

        lc = get_column_letter(next_col)
        formula1 = f"={_LOOKUPS_SHEET}!${lc}$1:${lc}${len(values)}"
        dv = DataValidation(type="list", formula1=formula1, allow_blank=True, showDropDown=False)
    else:
        quoted = ", ".join(f'"{v}"' for v in values)
        dv = DataValidation(
            type="list",
            formula1=f'"{",".join(values)}"',
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=True,
            errorTitle="Invalid value",
            error=f"Column {name} only accepts {quoted}.",
            showInputMessage=True,
        )

    dv.sqref = f"{col_letter}2:{col_letter}100"
    ws.add_data_validation(dv)


def create_template(
    columns: list[dict], output_path: str = "template.xlsx", max_col_width: int = 30
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    for col_idx, col_def in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_def["name"])
        col_letter = get_column_letter(col_idx)

        if col_def.get("values"):
            _add_validation(ws, wb, col_letter, col_def)

        candidates = [col_def["name"]] + (col_def.get("values") or [])
        max_width = col_def.get("width") or min(max(len(s) for s in candidates), max_col_width)
        ws.column_dimensions[col_letter].width = max_width
        for row in range(2, 101):
            ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True)

    wb.save(output_path)


if __name__ == "__main__":
    print(f"Template saved to: template.xlsx")
    create_template(
        columns=[
            {
                "name": "Species/Reference Text",
                "values": [
                    "Common coquí (Eleutherodactylus coqui) / Smith et al. (2017)",
                    "Red-breasted parakeet (Psittacula alexandri) / Neo et al. (2012)",
                    "House crow (Corvus splendens) / Roll et al. (2007)",
                    "Japanese quail (Coturnix japonica) / Barilani et al. (2005)",
                    "Common myna (Acridotheres tristis) / Hughes et al. (2017)",
                    "Asian common toad (Duttaphrynus melanostictus) / Church et al. (1960)",
                ],
            },
            {
                "name": "EICAT Category",
                "values": ["Massive", "Major", "Moderate", "Minor", "Minimal Concern", "Data Deficient", "No Alien Populations"],
            },
            {
                "name": "Impact mechanism",
                "values": [
                    "Competition", "Predation", "Hybridisation", "Transmission of disease",
                    "Parasitism", "Poisoning/toxicity", "Bio-fouling or other direct physical disturbance",
                    "Grazing/herbivory/browsing", "Chemical impact on ecosystem", "Physical impact on ecosystem",
                    "Structural impact on ecosystem", "Indirect impacts through interactions with other species",
                ],
            },
            {"name": "Evidence for EICAT impact category", "width": 60},
            {"name": "Confidence Rating", "values": ["low", "medium", "high"]},
            {"name": "Impacted native species"},
        ]
    )
