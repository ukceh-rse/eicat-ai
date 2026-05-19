"""Generate a references spreadsheet for researcher paper review.

Produces an xlsx with one row per unique reference containing:
  - Reference text
  - Clickable Google Scholar search link
  - Suggested PDF filename  (first_author_year.pdf)
  - EICAT categories covered by this reference (distinct, severity-ordered)
  - Impact mechanisms covered by this reference (distinct, frequency-ordered)
  - Empty "Available" column for the researcher to mark Y / N

Column headers have auto-filter dropdowns for sorting/filtering.

Usage:
    uv run utils/references.py
    uv run utils/references.py --input data/gisd_clean.csv --output data/references.xlsx
"""
import re
import urllib.parse
from pathlib import Path
from typing import Annotated

import pandas as pd
import typer
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

app = typer.Typer(add_completion=False)

DEFAULT_INPUT = Path(__file__).parent.parent / "data" / "gisd_clean.csv"
DEFAULT_OUTPUT = Path(__file__).parent.parent / "data" / "references.xlsx"


def first_author_year(reference: str) -> tuple[str, str]:
    surname = re.match(r"^([A-Za-z'][A-Za-z'-]+)", reference.strip())
    year = re.search(r"\b(19|20)\d{2}\b", reference)
    return (
        surname.group(1).lower() if surname else "unknown",
        year.group() if year else "unknown",
    )


def pdf_filename(reference: str) -> str:
    author, year = first_author_year(reference)
    return f"{author}_{year}.pdf"


def scholar_url(reference: str) -> str:
    return f"https://scholar.google.com/scholar?q={urllib.parse.quote_plus(reference[:200])}"


CAT_SEVERITY: dict[str, int] = {"MV": 5, "MR": 4, "MO": 3, "MN": 2, "MC": 1, "DD": 0}


def ref_summary(group: pd.DataFrame) -> dict:
    """Aggregate distinct categories and mechanisms for one reference."""
    cats = sorted(group["EICAT Category"].dropna().unique(), key=lambda c: -CAT_SEVERITY.get(c, 0))
    mech_counts = group["Impact mechanism"].dropna().value_counts()
    return {
        "categories": ", ".join(cats),
        "mechanisms": ", ".join(mech_counts.index.tolist()),
    }


@app.command()
def main(
    input: Annotated[Path, typer.Option(help="Cleaned GISD CSV")] = DEFAULT_INPUT,
    output: Annotated[Path, typer.Option(help="References xlsx output")] = DEFAULT_OUTPUT,
) -> None:
    """Create a references spreadsheet for paper availability review."""
    df = pd.read_csv(input)

    summaries = {
        ref: ref_summary(group)
        for ref, group in df.dropna(subset=["Reference"]).groupby("Reference", sort=False)
    }
    refs = sorted(summaries.keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "References"

    headers = ["Reference", "Scholar Link", "PDF Filename", "EICAT Categories", "Impact Mechanisms", "Available"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="366092")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 18

    link_font = Font(color="0563C1", underline="single")

    for ref in refs:
        row = ws.max_row + 1
        s = summaries[ref]

        ws.cell(row, 1, ref).alignment = Alignment(wrap_text=True, vertical="top")

        url = scholar_url(ref)
        link_cell = ws.cell(row, 2, "Scholar")
        link_cell.hyperlink = url
        link_cell.font = link_font
        link_cell.alignment = Alignment(horizontal="center", vertical="top")

        ws.cell(row, 3, pdf_filename(ref)).alignment = Alignment(vertical="top")
        ws.cell(row, 4, s["categories"]).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 5, s["mechanisms"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 6, "").alignment = Alignment(horizontal="center", vertical="top")

    # Y/N dropdown on Available column
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showDropDown=False)
    dv.sqref = f"F2:F{ws.max_row}"
    ws.add_data_validation(dv)

    # Auto-filter on all columns (enables sorting/filtering in Excel/LibreOffice)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 35
    ws.column_dimensions["F"].width = 12
    ws.freeze_panes = "A2"

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    typer.echo(f"References: {len(refs):,}")
    typer.echo(f"Saved:      {output}")


if __name__ == "__main__":
    app()
