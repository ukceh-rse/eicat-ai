"""Give the rendered PPTX talk the OneSTOP closing slide: ants photo, white text, a large EU emblem.

Pandoc picks layouts by name only, so the closing slide renders on the plain content layout. This
restyles it after `quarto render`, matching the template's closing page and the HTML version.

    uv run --with python-pptx docs/llmsig/finish_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).parent
DECK = ROOT.parent / "_site" / "llmsig" / "eicat-ai-talk.pptx"
BACKGROUND = ROOT / "assets" / "closing-ants.jpg"
EMBLEM = ROOT / "assets" / "eu-funded-white.png"  # white lettering, for dark backgrounds
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_text(slide, text: str, top: float, height: float, size: int, width: float = 11.0) -> None:
    left = (slide.part.package.presentation_part.presentation.slide_width - Inches(width)) // 2
    box = slide.shapes.add_textbox(Emu(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = WHITE


def main() -> None:
    deck = Presentation(DECK)
    slide = deck.slides[-1]
    assert slide.shapes.title.text == "Questions", "last slide is not the closing slide"
    width, height = deck.slide_width, deck.slide_height

    # Photo behind everything; it also covers the layout's logo and small emblem.
    photo = slide.shapes.add_picture(str(BACKGROUND), 0, 0, width, height)
    slide.shapes._spTree.remove(photo._element)
    slide.shapes._spTree.insert(2, photo._element)

    title = slide.shapes.title
    title.left, title.top, title.width, title.height = 0, Inches(0.5), width, Inches(1.2)
    for paragraph in title.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(54)
            run.font.bold = True
            run.font.color.rgb = WHITE

    # Replace the body ("Thank you" + funding text) with centred white text around a large emblem.
    body = next(s for s in slide.placeholders if s.placeholder_format.idx != 0)
    funding = body.text_frame.paragraphs[-1].text
    body._element.getparent().remove(body._element)
    add_text(slide, "Thank you", top=1.75, height=0.7, size=30)
    emblem_width = Inches(7)
    slide.shapes.add_picture(str(EMBLEM), (width - emblem_width) // 2, Inches(2.75), emblem_width)
    add_text(slide, funding, top=5.9, height=1.0, size=11)

    deck.save(DECK)
    print(f"wrote {DECK}")


if __name__ == "__main__":
    main()
