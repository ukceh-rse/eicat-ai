"""Turn the OneSTOP PowerPoint template into a Quarto/pandoc reference doc.

Pandoc picks slide layouts by name ("Title Slide", "Title and Content", ...). The OneSTOP
template (docs/assets/template.pptx, exported from Google Slides) names its layouts in
Bulgarian, so this renames the ones we use and drops the template's sample slides.

    uv run --with python-pptx docs/llmsig/make_reference.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).parent
TEMPLATE = ROOT.parent / "assets" / "template.pptx"
OUTPUT = ROOT / "assets" / "onestop-reference.pptx"

# template layout index -> pandoc layout name
LAYOUTS = {
    0: "Title Slide",  # leaf-ring title page
    1: "Title and Content",  # plain content page with logo and EU emblem
    4: "Section Header",  # full-bleed deer photo
    7: "Two Content",
    8: "Blank",  # title only
    # Unused by the talk, but must exist: pandoc writes its own fallback layouts over existing
    # slideLayoutN.xml files, clobbering the Section Header and Two Content layouts above.
    3: "Comparison",
    9: "Content with Caption",
}
# Pandoc fills content placeholders of type "obj" (or untyped); the Google Slides export types them "body".
CONTENT_LAYOUTS = ("Title and Content", "Two Content")
# Pandoc needs a slide-number field in the notes master to write speaker notes; the template has none.
SLIDE_NUMBER = f"""<p:sp {nsdecls("p", "a")}>
  <p:nvSpPr><p:cNvPr id="90" name="Slide Number Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
    <p:nvPr><p:ph type="sldNum" sz="quarter" idx="5"/></p:nvPr></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="3884613" y="8685213"/><a:ext cx="2971800" cy="458787"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
  <p:txBody><a:bodyPr anchor="b"/><a:lstStyle/><a:p><a:pPr algn="r"/>
    <a:fld id="{{18BDFEC3-8487-43E8-A154-7C12CBC1FFF2}}" type="slidenum"><a:rPr lang="en-GB"/><a:t>‹#›</a:t></a:fld>
  </a:p></p:txBody>
</p:sp>"""


def main() -> None:
    deck = Presentation(TEMPLATE)
    for index, name in LAYOUTS.items():
        layout = deck.slide_layouts[index]
        layout._element.cSld.set("name", name)
        if name in CONTENT_LAYOUTS:
            for placeholder in layout.placeholders:
                ph = placeholder._element.ph
                if ph.get("type") == "body":
                    del ph.attrib["type"]
    deck.notes_master._element.cSld.spTree.append(parse_xml(SLIDE_NUMBER))
    # Drop the template's embedded fonts: pandoc copies the font parts but not their content type,
    # which leaves PowerPoint offering to repair the file.
    fonts = deck.part._element.find(qn("p:embeddedFontLst"))
    if fonts is not None:
        for rel_id in {e.get(qn("r:id")) for e in fonts.iter() if e.get(qn("r:id"))}:
            deck.part.drop_rel(rel_id)
        deck.part._element.remove(fonts)
    slide_ids = deck.slides._sldIdLst
    for slide_id in list(slide_ids):
        deck.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    deck.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
